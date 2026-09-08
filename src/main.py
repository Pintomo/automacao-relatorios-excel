import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# =========================
# LER PLANILHA DE VENDAS
# =========================

df = pd.read_excel("data/vendas.xlsx")


# =========================
# CALCULAR INDICADORES
# =========================

faturamento_total = df["Valor"].sum()
quantidade_vendas = len(df)

if quantidade_vendas > 0:
    ticket_medio = faturamento_total / quantidade_vendas
else:
    ticket_medio = 0


# =========================
# CRIAR RELATÓRIO
# =========================

relatorio = pd.DataFrame({
    "Indicador": [
        "Faturamento Total",
        "Quantidade de Vendas",
        "Ticket Médio"
    ],
    "Valor": [
        faturamento_total,
        quantidade_vendas,
        ticket_medio
    ]
})


# Caminho do arquivo
caminho_relatorio = "reports/relatorio_vendas.xlsx"


# Criar arquivo Excel
relatorio.to_excel(
    caminho_relatorio,
    index=False,
    startrow=3
)


# =========================
# ABRIR PARA FORMATAR
# =========================

workbook = load_workbook(caminho_relatorio)
planilha = workbook.active

planilha.title = "Relatório de Vendas"


# =========================
# TÍTULO
# =========================

planilha.merge_cells("A1:B1")
planilha["A1"] = "RELATÓRIO DE VENDAS"

planilha["A1"].font = Font(
    bold=True,
    size=16,
    color="FFFFFF"
)

planilha["A1"].fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

planilha["A1"].alignment = Alignment(
    horizontal="center",
    vertical="center"
)

planilha.row_dimensions[1].height = 30


# =========================
# DATA E HORA
# =========================

agora = datetime.now()

planilha.merge_cells("A2:B2")
planilha["A2"] = f"Gerado em: {agora.strftime('%d/%m/%Y %H:%M')}"

planilha["A2"].font = Font(
    italic=True,
    size=10
)

planilha["A2"].alignment = Alignment(
    horizontal="center"
)


# =========================
# CABEÇALHO
# =========================

for celula in planilha[4]:

    celula.font = Font(
        bold=True,
        color="FFFFFF"
    )

    celula.fill = PatternFill(
        fill_type="solid",
        fgColor="4472C4"
    )

    celula.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


# =========================
# FORMATAR VALORES
# =========================

planilha["B5"].number_format = 'R$ #,##0.00'
planilha["B6"].number_format = '0'
planilha["B7"].number_format = 'R$ #,##0.00'


for linha in range(5, 8):
    planilha[f"B{linha}"].alignment = Alignment(
        horizontal="right"
    )


# =========================
# LARGURA DAS COLUNAS
# =========================

for numero_coluna in range(1, planilha.max_column + 1):

    letra_coluna = get_column_letter(numero_coluna)
    maior_tamanho = 0

    for numero_linha in range(1, planilha.max_row + 1):

        celula = planilha.cell(
            row=numero_linha,
            column=numero_coluna
        )

        if celula.value is not None:
            tamanho = len(str(celula.value))

            if tamanho > maior_tamanho:
                maior_tamanho = tamanho

    planilha.column_dimensions[letra_coluna].width = max(
        maior_tamanho + 5,
        20
    )


planilha.column_dimensions["A"].width = 25
planilha.column_dimensions["B"].width = 20


# =========================
# SALVAR
# =========================

workbook.save(caminho_relatorio)


# =========================
# RESULTADO NO TERMINAL
# =========================

print("=" * 40)
print("       RELATÓRIO DE VENDAS")
print("=" * 40)

print(f"Faturamento total: R$ {faturamento_total:.2f}")
print(f"Quantidade de vendas: {quantidade_vendas}")
print(f"Ticket médio: R$ {ticket_medio:.2f}")

print()
print("Relatório criado com sucesso!")
print(f"Arquivo: {caminho_relatorio}")